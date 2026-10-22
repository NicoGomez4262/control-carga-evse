# -*- coding: utf-8 -*-
"""F8: deja el proyecto listo para sacar los archivos de PCBWay (Altium CERRADO).

1. PcbDoc:
   - contorno de la placa en Mechanical 1 (4 pistas de 0,1 mm sobre el borde): los Gerber necesitan una capa con
     el contorno y el proyecto solo lo tenía en Board6. La capa V7 de Mechanical 1 es 01000201 (como 02000201 la 2).
   - BT1.3 (contacto de la pila, Ø 17 mm) sin pasta: expansión manual de pasta = −8,5 mm (bytes 86 y 101 del bloque
     de geometría del pad). Con la regla, el stencil le ponía pasta a todo el disco.
2. `Control de Carga IoT.OutJob` (estaba vacío) con, a partir de las plantillas de AD25 (OutputJobs\\Fabrication y
   Assembly):
   - Gerber RS-274X en pulgadas 2:5: Top/Bottom Layer, Overlay, Solder y Paste y Mechanical 1. Sin la capa
     Keep-Out, porque sus rellenos de la barrera se podrían leer como fresado.
   - NC Drill 2:5 con PTH/NPTH separados y ranuras G85 (las de J5).
   - Pick & Place CSV en mm, STEP y plano de ensamble (PDF).
3. docs/fabricacion/BOM_PCBWay.xlsx y .csv: a partir de bom_f7.LINEAS (la fuente única), con footprint y SMD/THT
   leídos de la PCB y agrupada por montaje (PCBWay / a mano / no se monta).

Uso: python docs/tools/salidas_f8.py [--escribir | --solo-outjob]
"""
import csv
import struct
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import altium_ole as ao  # noqa: E402
import aplicar_ruteo_f8 as ap  # noqa: E402
import pcblib  # noqa: E402
import revisar_pcb as rv  # noqa: E402

RAIZ = Path(__file__).resolve().parents[2]
PROY = rv.PCB.parent
OUTJOB = PROY / "Control de Carga IoT.OutJob"
FAB = RAIZ / "docs" / "fabricacion"
U, O = rv.U, 25.4
MECH1, MECH1_V7 = 57, bytes.fromhex("01000201")
LARGO, ALTO = 130.0, 100.0
SALIDA = r"Project Outputs for Control de Carga IoT\PCBWay"

# ---------------------------------------------------------------- PcbDoc
def contorno_mech1(tr):
    """Pistas libres de Mechanical 1 que ya estén (para no duplicar) y las 4 que faltan."""
    ya = [bl[0] for t, bl in tr if bl[0][0] == MECH1 and struct.unpack_from("<H", bl[0], 7)[0] == 0xFFFF]
    if ya:
        return []
    plantilla = next(bl[0] for t, bl in tr if struct.unpack_from("<H", bl[0], 7)[0] == 0xFFFF and bl[0][0] == 1)
    pts = [(0, 0), (LARGO, 0), (LARGO, ALTO), (0, ALTO), (0, 0)]
    nuevas = []
    for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
        g = bytearray(ap.pista_bin(plantilla, 0xFFFF, 1, x1, y1, x2, y2, 0.1))
        g[0] = MECH1
        g[41:45] = MECH1_V7
        nuevas.append((4, [bytes(g)]))
    return nuevas


def sin_pasta_bt1(pads, orden):
    cambios = 0
    for n, (t, bl) in enumerate(pads):
        g = bl[4]
        comp = struct.unpack_from("<H", g, 7)[0]
        nombre = bl[0][1:1 + bl[0][0]].decode("latin-1")
        if comp == 0xFFFF or orden[comp] != "BT1" or nombre != "3":
            continue
        sx = struct.unpack_from("<i", g, 21)[0] * U
        g = bytearray(g)
        struct.pack_into("<i", g, 86, int(round(-sx / 2 / U)))
        g[101] = 2
        bl = list(bl)
        bl[4] = bytes(g)
        pads[n] = (t, bl)
        cambios += 1
    return cambios


# ---------------------------------------------------------------- OutJob
CAPAS_GERBER = [16777217, 16842751, 16973830, 16973831, 16973832, 16973833, 16973834, 16973835, 16908289]


def outjob():
    plot = ",".join(f"{c}~1" for c in CAPAS_GERBER)
    gerber = ("AddToAllLayerClasses.Set= |AddToAllPlots.Set=SerializeLayerHash.Version~2,ClassName~TLayerToBoolean|"
              "CentrePlots=False|DrillDrawingSymbol=GraphicsSymbol|DrillDrawingSymbolSize=500000|EmbeddedApertures=True|"
              "FilmBorderSize=10000000|FilmXSize=200000000|FilmYSize=160000000|FlashAllFills=False|FlashPadShapes=True|"
              "G54OnApertureChange=False|GenerateDRCRulesFile=False|GenerateReliefShapes=True|GerberUnit=Imperial|"
              "IncludeUnconnectedMidLayerPads=False|LayerClassesMirror.Set= |LayerClassesPlot.Set= |"
              "LeadingAndTrailingZeroesMode=SuppressLeadingZeroes|MaxApertureSize=2500000|MinusApertureTolerance=50|"
              "Mirror.Set=SerializeLayerHash.Version~2,ClassName~TLayerToBoolean|MirrorDrillDrawingPlots=False|"
              "MirrorDrillGuidePlots=False|NumberOfDecimals=5|OptimizeChangeLocationCommands=True|OriginPosition=Relative|"
              f"Panelize=False|Plot.Set=SerializeLayerHash.Version~2,ClassName~TLayerToBoolean,{plot}|"
              "PlotPositivePlaneLayers=False|PlotUsedDrillDrawingLayerPairs=False|PlotUsedDrillGuideLayerPairs=False|"
              "PlusApertureTolerance=50|Record=GerberView|SoftwareArcs=False|Sorted=False")
    taladro = ("BoardEdgeRoutToolDia=2000000|GenerateBoardEdgeRout=False|GenerateDrilledSlotsG85=True|"
               "GenerateEIADrillFile=False|GenerateSeparatePlatedNonPlatedFiles=True|NumberOfDecimals=5|NumberOfUnits=2|"
               "OptimizeChangeLocationCommands=True|OriginPosition=Relative|Record=DrillView|Units=Imperial|"
               "ZeroesMode=SuppressTrailingZeroes")
    cols = [("Designator", False, True), ("Comment", False, True), ("Layer", False, True), ("Footprint", False, True),
            ("Center-X", True, True), ("Center-Y", True, True), ("Rotation", False, True), ("Description", False, True),
            ("ComponentKind", False, False), ("Footprint Description", False, False), ("Height", True, False),
            ("Pad-X", True, False), ("Pad-Y", True, False), ("Ref-X", True, False), ("Ref-Y", True, False),
            ("Variation", False, False)]
    pp = ("Record=PickPlaceView|Units=Metric|GenerateCSVFormat=True|GenerateTextFormat=False|ShowUnits=False|Separator=.|"
          "ExcludeFilterParam=False|IncludeVariations=False|Filter= |FilterActive=False|" +
          "|".join(f"Column#{i}=Name:{n},Fixed:True,Metric:{m},Visible:{v},Sort:None,Position:{i - 1}"
                   for i, (n, m, v) in enumerate(cols, 1)))
    step = ("Record=ExportSTEPView|ExportComponentOptions=0|ExportModelsOption=2|ExportHolesOption=0|"
            "CanSelectPrimitives=False|IncludeMechanicalPadHoles=False|IncludeElectricalPadHoles=False|"
            "IncludeFreePadHoles=False|ExportFoldedBoard=True|ExportFoldedBoardRate=100|ComponentSuffixType=0|"
            "ComponentSuffix= |ExportCopperOption=0|ExportCopperLayer=1|ExportPadAndViaBarrelsOnly=False|"
            "IgnoreBoardCopperLayerColors=False|ExportAsSinglePart=False")
    capa = ("CArc=Full|CFill=Full|Comment=Full|Coordinate=Full|CPad=Full|CRegion=Full|CText=Full|CTrack=Full|CVia=Full|"
            "Designator=Full|Dimension=Full|DLayer1=TopLayer|DLayer2=BottomLayer|FArc=Full|FFill=Full|FPad=Full|"
            "FRegion=Full|FText=Full|FTrack=Full|FVia=Full|Layer={}|Polygon=Full|PrintOutIndex=0|Record=PcbPrintLayer")
    # en el plano de ensamble, de Top Layer solo los pads (sin pistas, vías, gotas ni el polígono de GND)
    solo_pads = ("CArc=Hidden|CFill=Hidden|Comment=Full|Coordinate=Full|CPad=Full|CRegion=Hidden|CText=Hidden|"
                 "CTrack=Hidden|CVia=Hidden|Designator=Full|Dimension=Full|DLayer1=TopLayer|DLayer2=BottomLayer|"
                 "FArc=Hidden|FFill=Hidden|FPad=Full|FRegion=Hidden|FText=Hidden|FTrack=Hidden|FVia=Hidden|"
                 "Layer=TopLayer|Polygon=Hidden|PrintOutIndex=0|Record=PcbPrintLayer")
    salidas = [  # (tipo, nombre, categoría, medio, orden en el medio, configuraciones, opciones de página)
        ("Gerber", "Gerber", "Fabrication", 3, 1, [gerber], None),
        ("NC Drill", "NC Drill", "Fabrication", 3, 2, [taladro], None),
        ("Pick Place", "Pick and Place", "Assembly", 4, 1, [pp], None),
        ("ExportSTEP", "STEP", "Export", 4, 2, [step], None),
        ("Assembly", "Plano de ensamble", "Assembly", 2, 1,
         ["DesignatorDisplayMode=Physical|PrintArea=DesignExtent|PrintAreaLowerLeftCornerX=0|PrintAreaLowerLeftCornerY=0|"
          "PrintAreaUpperRightCornerX=0|PrintAreaUpperRightCornerY=0|Record=PcbPrintView",
          "IncludeBottomLayerComponents=False|IncludeMultiLayerComponents=True|IncludeTopLayerComponents=True|"
          "IncludeViewports=True|Index=0|Mirror=False|Name=Ensamble TOP|PadNumberFontSize=14|Record=PcbPrintOut|"
          "ShowHoles=False|ShowPadNets=False|ShowPadNumbers=False|SubstituteFonts=False",
          solo_pads, capa.format("TopOverlay"), capa.format("Mechanical1")],
         "Record=PageOptions|CenterHorizontal=True|CenterVertical=True|PrintScale=1.00|XCorrection=1.00|"
         "YCorrection=1.00|PrintKind=1|BorderSize=5000000|LeftOffset=0|BottomOffset=0|Orientation=2|PaperLength=1000|"
         "PaperWidth=1000|Scale=100|PaperSource=7|PrintQuality=-3|MediaType=1|DitherType=10|PrintScaleMode=1|"
         "PaperKind=Letter|PaperIndex=1"),
    ]
    L = ["[OutputJobFile]", "Version=1.0"] + [f"{k}=" for k in (
        "Caption", "Description", "VaultGUID", "ItemGUID", "ItemHRID", "RevisionGUID", "RevisionId", "VaultHRID",
        "AutoItemHRID", "NextRevId", "FolderGUID", "LifeCycleDefinitionGUID", "RevisionNamingSchemeGUID")] + [""]
    impresora = "Record=PrinterOptions|Copies=1|Duplex=1|TrueTypeOptions=3|Collate=1|PrintJobKind=1|PrintWhat=1"
    L += ["[OutputGroup1]", "Name=PCBWay F8", "Description=Gerber, taladros, Pick and Place, STEP y plano de ensamble",
          "TargetOutputMedium=Gerber y taladros", "VariantName=[No Variations]", "VariantScope=0",
          "CurrentConfigurationName=", "TargetPrinter=Microsoft Print to PDF", f"PrinterOptions={impresora}",
          "OutputMedium1=Print Job", "OutputMedium1_Type=Printer", "OutputMedium1_Printer=",
          f"OutputMedium1_PrinterOptions={impresora}",
          "OutputMedium2=PDF", "OutputMedium2_Type=Publish",
          "OutputMedium3=Gerber y taladros", "OutputMedium3_Type=GeneratedFiles",
          "OutputMedium4=Ensamble", "OutputMedium4_Type=GeneratedFiles"]
    for i, (tipo, nombre, cat, medio, pos, confs, pagina) in enumerate(salidas, 1):
        L += [f"OutputType{i}={tipo}", f"OutputName{i}={nombre}", f"OutputCategory{i}={cat}",
              f"OutputDocumentPath{i}=EVSE_Monitor.PcbDoc", f"OutputVariantName{i}=", f"OutputEnabled{i}=1"]
        L += [f"OutputEnabled{i}_OutputMedium{m}={pos if m == medio else 0}" for m in (1, 2, 3, 4)]
        L += [f"OutputDefault{i}=0"]
        if pagina:
            L += [f"PageOptions{i}={pagina}"]
        for k, c in enumerate(confs, 1):
            L += [f"Configuration{i}_Name{k}=OutputConfigurationParameter{k}", f"Configuration{i}_Item{k}={c}"]
        if tipo == "Assembly":
            L += [f"PcbPrintPreferences{i}="]
    pdf = SALIDA + r"\Ensamble_EVSE_Monitor.pdf"
    L += ["", "[PublishSettings]", f"OutputFilePath2={pdf}", "ReleaseManaged2=1",
          r"OutputBasePath2=Project Outputs for Control de Carga IoT", "OutputPathMedia2=", "OutputPathMediaValue2=",
          "OutputPathOutputer2=", "OutputPathOutputerPrefix2=", "OutputPathOutputerValue2=",
          "OutputFileName2=Ensamble_EVSE_Monitor.pdf", "OutputFileNameMulti2=", "UseOutputNameForMulti2=0",
          "OutputFileNameSpecial2=", "OpenOutput2=1", "PromptOverwrite2=0", "PublishMethod2=0", "ZoomLevel2=50",
          "FitSCHPrintSizeToDoc2=1", "FitPCBPrintSizeToDoc2=1", "GenerateNetsInfo2=1", "MarkPins2=1",
          "MarkNetLabels2=1", "MarkPortsId2=1", "GenerateTOC2=1", "ShowComponentParameters2=1", "GlobalBookmarks2=0",
          "PDFACompliance2=Disabled", "PDFVersion2=Default"]
    for m, sub in ((3, "Gerber"), (4, "Ensamble")):
        ruta = SALIDA + "\\" + sub + "\\"
        L += [f"OutputFilePath{m}={ruta}", f"ReleaseManaged{m}=1", f"OutputBasePath{m}={SALIDA}\\{sub}",
              f"OutputPathMedia{m}=", f"OutputPathMediaValue{m}=", f"OutputPathOutputer{m}=",
              f"OutputPathOutputerPrefix{m}=", f"OutputPathOutputerValue{m}=", f"OutputFileName{m}=",
              f"OutputFileNameMulti{m}=", f"UseOutputNameForMulti{m}=1", f"OutputFileNameSpecial{m}=",
              f"OpenOutput{m}=1"]
    L += ["", "[GeneratedFilesSettings]", f"RelativeOutputPath2={pdf}", "OpenOutputs2=1"]
    for m, sub in ((3, "Gerber"), (4, "Ensamble")):
        L += [f"RelativeOutputPath{m}={SALIDA}\\{sub}\\", f"OpenOutputs{m}=1", f"AddToProject{m}=0",
              f"TimestampFolder{m}=0", f"UseOutputName{m}=0", f"OpenODBOutput{m}=0", f"OpenGerberOutput{m}=0",
              f"OpenNCDrillOutput{m}=0", f"OpenIPCOutput{m}=0", f"EnableReload{m}=0"]
    return "\r\n".join(L) + "\r\n"


# ---------------------------------------------------------------- BOM de PCBWay
def bom():
    import bom_f7
    comps = rv.texto("Components6/Data")
    info = {ao.valor(c, "SOURCEDESIGNATOR"): dict(fp=ao.valor(c, "PATTERN"), desc=ao.valor(c, "SOURCEDESCRIPTION") or "")
            for c in comps}
    orden = [ao.valor(c, "SOURCEDESIGNATOR") for c in comps]
    tht = defaultdict(bool)
    for p in rv.pads():
        if p["comp"] != 0xFFFF and p["hueco"] > 0:
            tht[orden[p["comp"]]] = True
    filas, vistos = [], set()
    for ln in bom_f7.LINEAS:
        des = [d.strip() for d in ln[0].split(",")]
        if not all(d in info for d in des):
            continue
        vistos.update(des)
        fps = sorted({info[d]["fp"] for d in des})
        tipo = "THT" if any(tht[d] for d in des) else "SMD"
        montaje = {"PCBWay": "PCBWay", "Manual": "A mano (no ensamblar)", "DNP": "No se monta (DNP)"}[ln[5]]
        filas.append(dict(montaje=montaje, designadores=", ".join(des), cantidad=len(des), fabricante=ln[1], mpn=ln[2],
                          descripcion=info[des[0]]["desc"], footprint=" / ".join(fps), tipo=tipo, proveedor=ln[3],
                          pn_proveedor=ln[4], nota=ln[7]))
    faltan = sorted(set(info) - vistos)
    orden_m = {"PCBWay": 0, "A mano (no ensamblar)": 1, "No se monta (DNP)": 2}
    filas.sort(key=lambda f: (orden_m[f["montaje"]], f["designadores"]))
    return filas, faltan


def escribir_bom(filas):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    FAB.mkdir(parents=True, exist_ok=True)
    cab = ["Item #", "Designator", "Qty", "Manufacturer", "Mfg Part #", "Description", "Package / Footprint", "Type",
           "Assembly", "Supplier", "Supplier PN", "Notes"]
    claves = ["designadores", "cantidad", "fabricante", "mpn", "descripcion", "footprint", "tipo", "montaje",
              "proveedor", "pn_proveedor", "nota"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM PCBWay"
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for n, f in enumerate(filas, 1):
        ws.append([n] + [f[k] for k in claves])
        if f["montaje"] != "PCBWay":
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
    for col, w in zip("ABCDEFGHIJKL", (7, 34, 5, 24, 26, 40, 34, 6, 22, 11, 16, 40)):
        ws.column_dimensions[col].width = w
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(FAB / "BOM_PCBWay.xlsx")
    with open(FAB / "BOM_PCBWay.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(cab)
        for n, f in enumerate(filas, 1):
            w.writerow([n] + [f[k] for k in claves])


def main():
    filas, faltan = bom()
    por = defaultdict(int)
    for f in filas:
        por[f["montaje"]] += f["cantidad"]
    print(f"BOM: {len(filas)} líneas, {sum(f['cantidad'] for f in filas)} piezas {dict(por)}; sin línea: {faltan}")
    tr = ap.registros(ao.leer_stream(rv.PCB, "Tracks6/Data"), 1)
    nuevas = contorno_mech1(tr)
    pads = ap.registros(ao.leer_stream(rv.PCB, "Pads6/Data"), 6)
    orden = [ao.valor(c, "SOURCEDESIGNATOR") for c in rv.texto("Components6/Data")]
    n_bt1 = sin_pasta_bt1(pads, orden)
    print(f"Contorno en Mechanical 1: {len(nuevas)} pistas nuevas; BT1.3 sin pasta: {n_bt1} pad")
    if "--solo-outjob" in sys.argv:          # 22 oct: rehacer solo el OutJob (el PcbDoc ya tiene los cambios)
        OUTJOB.write_bytes(outjob().encode("utf-8"))
        print("OutJob escrito:", OUTJOB)
        return
    if "--escribir" not in sys.argv:
        print("(prueba: usa --escribir)")
        return
    if ao.altium_abierto():
        sys.exit("Altium está abierto (X2.exe): guarda, ciérralo y vuelve a correr.")
    print("Respaldo:", ao.respaldo(rv.PCB, "antes_salidas_f8"))
    tr += nuevas
    ao.escribir_streams(rv.PCB, {"Tracks6/Data": ap.unir(tr), "Tracks6/Header": struct.pack("<I", len(tr)),
                                 "Pads6/Data": ap.unir(pads)})
    if OUTJOB.exists():
        print("Respaldo:", ao.respaldo(OUTJOB, "antes_salidas_f8"))
    OUTJOB.write_bytes(outjob().encode("utf-8"))
    escribir_bom(filas)
    print("Escrito: PcbDoc, OutJob y", FAB / "BOM_PCBWay.xlsx")


if __name__ == "__main__":
    main()
